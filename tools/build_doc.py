#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
功能实现.md  ->  功能实现.docx

用法（在 E:\\Java-demo\\Cinema 目录下）：
    python tools\\build_doc.py
    python tools\\build_doc.py 功能实现.md 功能实现.docx

说明：
  * 唯一的源文件是 功能实现.md，改内容只改它，然后重跑本脚本覆盖 docx。
  * 支持的 Markdown 子集：# ~ #### 标题、``` 代码块、表格、> 引用、- / * 无序列表、
    1. 有序列表、--- 分隔线、`行内代码`、**加粗**、[文字](链接)（链接只保留文字）。
"""
import os
import re
import sys

try:
    from docx import Document
except ImportError:  # python-docx 没装
    print("缺少 python-docx，请用装了 python-docx 的解释器运行本脚本。")
    print("本机可用：C:\\Users\\15181\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe")
    print("或者直接双击 tools\\build_doc.bat")
    sys.exit(1)

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
CODE_FONT = "Consolas"

CODE_BG = "F6F8FA"
CODE_BORDER = "D0D7DE"
QUOTE_BORDER = "9DB7DD"
RULE_COLOR = "D9D9D9"

TEXT_WIDTH_CM = 16.4


# ---------------------------------------------------------------- 基础工具
def set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _border(par, edge, color, sz="6"):
    pPr = par._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    pbdr.append(el)


def left_border(par, color, sz="18"):
    _border(par, "left", color, sz)


def bottom_border(par, color):
    _border(par, "bottom", color)


def strip_links(text):
    return re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)


def add_inline(par, text, base_size=10.5, base_color=None):
    """把 **加粗** 和 `行内代码` 转换成 run"""
    text = strip_links(text)
    pos = 0
    for m in re.finditer(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            set_font(r, BODY_FONT, base_size, color=base_color)
        token = m.group(0)
        if token.startswith("**"):
            r = par.add_run(token[2:-2])
            set_font(r, BODY_FONT, base_size, bold=True, color=base_color)
        else:
            r = par.add_run(token[1:-1])
            set_font(r, CODE_FONT, base_size - 0.5, color=RGBColor(0xB0, 0x30, 0x60))
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        set_font(r, BODY_FONT, base_size, color=base_color)


def style_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(4)

    specs = {
        "Heading 1": (17, RGBColor(0x1F, 0x38, 0x64), 16, 8),
        "Heading 2": (13.5, RGBColor(0x2E, 0x54, 0x96), 12, 6),
        "Heading 3": (11, RGBColor(0x40, 0x40, 0x40), 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.2

    # 页脚页码
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    set_font(run, BODY_FONT, 9, color=RGBColor(0x80, 0x80, 0x80))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


# ---------------------------------------------------------------- 表格
def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def disp(s):
        s = strip_links(s)
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    weights = []
    for ci in range(cols):
        longest = max([disp(r[ci]) for r in rows if ci < len(r)] or [4])
        weights.append(min(max(longest, 4), 46))
    total = float(sum(weights))
    widths = [max(1.5, TEXT_WIDTH_CM * w / total) for w in weights]
    scale = TEXT_WIDTH_CM / sum(widths)
    widths = [w * scale for w in widths]

    for ri, row in enumerate(rows):
        trPr = table.rows[ri]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.1
            text = row[ci] if ci < len(row) else ""
            add_inline(par, text, base_size=8.5)
            if ri == 0:
                for r in par.runs:
                    r.bold = True
    doc.add_paragraph()


# ---------------------------------------------------------------- 主流程
def build(src, dst):
    with open(src, encoding="utf-8") as f:
        md = f.read()

    doc = Document()
    style_doc(doc)

    lines = md.split("\n")
    i = 0
    in_code = False
    code_buf = []
    table_buf = []
    last_kind = None
    num = 0

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        s = lines[i]
        stripped = s.strip()

        # 代码块
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                par = doc.add_paragraph()
                pf = par.paragraph_format
                pf.left_indent = Cm(0.25)
                pf.right_indent = Cm(0.1)
                pf.space_before = Pt(3)
                pf.space_after = Pt(5)
                pf.line_spacing = 1.0
                pf.keep_together = False
                shade(par, CODE_BG)
                left_border(par, CODE_BORDER)
                for idx, code_line in enumerate(code_buf):
                    run = par.add_run(code_line)
                    set_font(run, CODE_FONT, 8.5, color=RGBColor(0x24, 0x29, 0x2F))
                    if idx < len(code_buf) - 1:
                        run.add_break()
                last_kind = "code"
            i += 1
            continue

        if in_code:
            code_buf.append(s)
            i += 1
            continue

        # 表格
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            table_buf.append(cells)
            last_kind = "table"
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        # 分隔线
        if stripped == "---":
            if last_kind != "rule":
                par = doc.add_paragraph()
                par.paragraph_format.space_before = Pt(2)
                par.paragraph_format.space_after = Pt(6)
                bottom_border(par, RULE_COLOR)
            last_kind = "rule"
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = strip_links(m.group(2)).replace("`", "").strip()
            par = doc.add_heading(level=min(level, 3))
            par.text = ""
            add_inline(par, title, base_size={1: 17, 2: 13.5, 3: 11}.get(min(level, 3), 10.5))
            color = {1: RGBColor(0x1F, 0x38, 0x64),
                     2: RGBColor(0x2E, 0x54, 0x96),
                     3: RGBColor(0x40, 0x40, 0x40)}.get(min(level, 3))
            for r in par.runs:
                set_font(r, HEAD_FONT, None, bold=True, color=color)
            last_kind = "heading"
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.4)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, stripped[1:].strip(), base_color=RGBColor(0x44, 0x44, 0x44))
            left_border(par, QUOTE_BORDER)
            last_kind = "quote"
            i += 1
            continue

        # 有序列表（手写编号，避免 Word 跨列表连续编号）
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            if last_kind != "num":
                num = int(m.group(1))
            else:
                num += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.85)
            par.paragraph_format.first_line_indent = Cm(-0.55)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, str(num) + ".  " + m.group(2))
            last_kind = "num"
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, m.group(1))
            last_kind = "bullet"
            i += 1
            continue

        par = doc.add_paragraph()
        add_inline(par, stripped)
        last_kind = "para"
        i += 1

    flush_table()
    doc.save(dst)
    return doc


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(here)
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(root, "功能实现.md")
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(root, "功能实现.docx")
    doc = build(src, dst)
    print("md  :", src)
    print("docx:", dst)
    print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))


if __name__ == "__main__":
    main()